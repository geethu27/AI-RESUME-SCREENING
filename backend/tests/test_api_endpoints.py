import os
import sys
import unittest
import io
import docx

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app import create_app
from models.database import get_db


class TestAPIEndpoints(unittest.TestCase):

    def setUp(self):
        self.app = create_app()
        self.client = self.app.test_client()

    def test_health_endpoint(self):
        """Test GET /api/health."""
        response = self.client.get("/api/health")
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertTrue(data.get("success"))
        self.assertEqual(data.get("message"), "Resume Screening backend is running")

    def test_auth_login(self):
        """Test POST /api/login."""
        # Success with default admin
        res = self.client.post("/api/login", json={
            "email": "admin@company.com",
            "password": "admin123"
        })
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.get_json()["success"])

        # Invalid password
        bad_res = self.client.post("/api/login", json={
            "email": "admin@company.com",
            "password": "wrongpassword"
        })
        self.assertEqual(bad_res.status_code, 401)
        self.assertFalse(bad_res.get_json()["success"])

    def test_jobs_crud(self):
        """Test Jobs CRUD operations."""
        # Create Job
        payload = {
            "title": "Full Stack Engineer",
            "description": "Looking for React and Node.js developers",
            "required_skills": "React, Node.js, JavaScript",
            "preferred_skills": "Docker, Git",
            "minimum_experience": 2,
            "education": "B.Tech"
        }
        res = self.client.post("/api/jobs", json=payload)
        self.assertEqual(res.status_code, 201)
        job_id = res.get_json()["data"]["id"]

        # Get all jobs
        list_res = self.client.get("/api/jobs")
        self.assertEqual(list_res.status_code, 200)
        self.assertTrue(any(j["id"] == job_id for j in list_res.get_json()["data"]))

        # Get single job
        get_res = self.client.get(f"/api/jobs/{job_id}")
        self.assertEqual(get_res.status_code, 200)
        self.assertEqual(get_res.get_json()["data"]["title"], "Full Stack Engineer")

        # Update job
        update_res = self.client.put(f"/api/jobs/{job_id}", json={
            "title": "Senior Full Stack Engineer",
            "description": "Updated description",
            "required_skills": "React, Node.js, TypeScript",
            "preferred_skills": "Docker, Kubernetes",
            "minimum_experience": 4,
            "education": "B.Tech / M.Tech"
        })
        self.assertEqual(update_res.status_code, 200)
        self.assertEqual(update_res.get_json()["data"]["title"], "Senior Full Stack Engineer")

        # Delete job
        del_res = self.client.delete(f"/api/jobs/{job_id}")
        self.assertEqual(del_res.status_code, 200)

        # Confirm deleted
        check_res = self.client.get(f"/api/jobs/{job_id}")
        self.assertEqual(check_res.status_code, 404)

    def test_resume_upload_and_screening(self):
        """Test complete workflow: create job, upload resume, screen, and get results."""
        # 1. Create Job
        job_res = self.client.post("/api/jobs", json={
            "title": "Python Backend Engineer",
            "description": "Python, Flask, PostgreSQL engineer needed",
            "required_skills": "Python, Flask, PostgreSQL",
            "preferred_skills": "Docker, Git",
            "minimum_experience": 1,
            "education": "B.Tech"
        })
        job_id = job_res.get_json()["data"]["id"]

        # 2. Create in-memory test DOCX resume
        doc = docx.Document()
        doc.add_heading("John Candidate", 0)
        doc.add_paragraph("john.test@example.com | +91 9988776655")
        doc.add_paragraph("Education: B.Tech in Computer Science")
        doc.add_paragraph("Worked as a Python developer for 3 years with Flask, PostgreSQL and Docker.")
        docx_io = io.BytesIO()
        doc.save(docx_io)
        docx_io.seek(0)

        # 3. Upload resume
        upload_res = self.client.post(
            "/api/resumes/upload",
            data={"resumes": (docx_io, "john_candidate.docx")},
            content_type="multipart/form-data"
        )
        self.assertEqual(upload_res.status_code, 201)
        self.assertTrue(upload_res.get_json()["success"])

        # 4. Trigger Screening
        screen_res = self.client.post(f"/api/screening/{job_id}")
        self.assertEqual(screen_res.status_code, 200)
        results = screen_res.get_json()["data"]["results"]
        self.assertTrue(len(results) >= 1)

        # 5. Verify screening results GET
        get_screen_res = self.client.get(f"/api/screening/{job_id}")
        self.assertEqual(get_screen_res.status_code, 200)

        # Clean up job
        self.client.delete(f"/api/jobs/{job_id}")


if __name__ == "__main__":
    unittest.main()
