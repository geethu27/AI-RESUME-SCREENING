import os
import sys
import unittest
import tempfile
import docx
from reportlab.pdfgen import canvas

# Ensure backend directory is in python path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from services.resume_parser import extract_text_from_pdf, extract_text_from_docx, parse_resume
from services.text_preprocessor import clean_and_preprocess
from services.skill_extractor import (
    extract_skills_from_text,
    extract_job_skills,
    extract_experience,
    extract_education,
    extract_contact_info,
    DEFAULT_SKILL_DICTIONARY
)
from services.matching_engine import match_resume_with_job
from services.scoring_engine import calculate_matching_score
from services.ranking_engine import rank_candidates


class TestResumeScreeningServices(unittest.TestCase):

    def test_text_preprocessing(self):
        """Test text cleaning and special term preservation."""
        input_text = "Experienced Python Developer with 2+ years of experience in Flask, React and MySQL."
        cleaned = clean_and_preprocess(input_text)
        self.assertIn("python", cleaned)
        self.assertIn("flask", cleaned)
        self.assertIn("react", cleaned)
        self.assertIn("mysql", cleaned)
        self.assertIn("2 years", cleaned)

        # Test technical term preservation
        cplus_text = "Skills include C++, C#, .NET and Node.js."
        cplus_cleaned = clean_and_preprocess(cplus_text)
        self.assertIn("c++", cplus_cleaned)
        self.assertIn("c#", cplus_cleaned)
        self.assertIn(".net", cplus_cleaned)
        self.assertIn("node.js", cplus_cleaned)

    def test_docx_extraction(self):
        """Test DOCX resume text extraction."""
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            doc = docx.Document()
            doc.add_paragraph("Alice Johnson")
            doc.add_paragraph("Python and Django Developer")
            doc.save(tmp.name)
            tmp_path = tmp.name

        try:
            text, success, err = extract_text_from_docx(tmp_path)
            self.assertTrue(success)
            self.assertIn("Alice Johnson", text)
            self.assertIn("Python and Django Developer", text)
            self.assertEqual(err, "")
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_pdf_extraction(self):
        """Test PDF resume text extraction using ReportLab and PyPDF2."""
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            c = canvas.Canvas(tmp.name)
            c.drawString(100, 750, "Bob Smith")
            c.drawString(100, 730, "Software Engineer with React and Docker experience")
            c.save()
            tmp_path = tmp.name

        try:
            text, success, err = extract_text_from_pdf(tmp_path)
            self.assertTrue(success)
            self.assertIn("Bob Smith", text)
            self.assertIn("React and Docker", text)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_skill_extraction(self):
        """Test keyword-based skill extraction."""
        resume_text = "I am a Python developer with Flask, MySQL, Docker and AWS experience."
        skills = extract_skills_from_text(resume_text)
        self.assertIn("python", skills)
        self.assertIn("flask", skills)
        self.assertIn("mysql", skills)
        self.assertIn("docker", skills)
        self.assertIn("aws", skills)
        self.assertNotIn("java", skills)

    def test_job_skill_extraction(self):
        """Test splitting and deduplication of required and preferred skills."""
        req_str = "Python, Flask, React, MySQL"
        pref_str = "Git, REST API, Docker, Python"  # Python is duplicated
        res = extract_job_skills(req_str, pref_str)
        self.assertEqual(res["required_skills"], ["python", "flask", "react", "mysql"])
        self.assertIn("git", res["preferred_skills"])
        self.assertIn("docker", res["preferred_skills"])
        self.assertIn("rest api", res["preferred_skills"])
        self.assertNotIn("python", res["preferred_skills"])  # deduplicated from preferred

    def test_experience_extraction(self):
        """Test regex-based experience extraction."""
        text1 = "Worked as a Python Developer for 2 years."
        self.assertEqual(extract_experience(text1), 2.0)

        text2 = "Senior Architect with 5+ years of experience in cloud computing."
        self.assertEqual(extract_experience(text2), 5.0)

        text3 = "No prior experience stated in this document."
        self.assertEqual(extract_experience(text3), 0.0)

    def test_education_extraction(self):
        """Test education degree and branch extraction."""
        text1 = "Completed B.Tech in Computer Science Engineering in 2022."
        edu1 = extract_education(text1)
        self.assertEqual(edu1["degree"], "B.Tech")
        self.assertIn("Computer Science", edu1["summary"])

        text2 = "Holds an MCA degree from Delhi University."
        edu2 = extract_education(text2)
        self.assertEqual(edu2["degree"], "MCA")

    def test_contact_info_extraction(self):
        """Test email, phone, and name heuristics."""
        text = "Jane Doe\njane.doe@example.com\n+91 9876543210\nSoftware Developer"
        contact = extract_contact_info(text)
        self.assertEqual(contact["name"], "Jane Doe")
        self.assertEqual(contact["email"], "jane.doe@example.com")
        self.assertIn("9876543210", contact["phone"])

    def test_rule_based_matching_and_scoring_section_16_example(self):
        """
        Verify exact Section 16 specification example:
        Job: Python, Flask, React, MySQL (Required), Git, Docker (Preferred), Exp: 1, Edu: B.Tech
        Resume: Python, Flask, MySQL, Git, Exp: 1, Edu: B.Tech
        Score: 3/4*60 (45) + 1/2*15 (7.5) + 15 + 10 = 77.5 => Review
        """
        req_skills = ["python", "flask", "react", "mysql"]
        pref_skills = ["git", "docker"]
        cand_skills = ["python", "flask", "mysql", "git"]

        match_res = match_resume_with_job(req_skills, pref_skills, cand_skills)
        self.assertEqual(match_res["matched_required_count"], 3)
        self.assertEqual(match_res["total_required"], 4)
        self.assertEqual(match_res["matched_preferred_count"], 1)
        self.assertEqual(match_res["total_preferred"], 2)

        score_res = calculate_matching_score(
            matching_result=match_res,
            candidate_experience=1.0,
            required_experience=1.0,
            candidate_education="B.Tech in Computer Science",
            required_education="B.Tech"
        )

        self.assertEqual(score_res["required_score"], 45.0)
        self.assertEqual(score_res["preferred_score"], 7.5)
        self.assertEqual(score_res["experience_score"], 15.0)
        self.assertEqual(score_res["education_score"], 10.0)
        self.assertEqual(score_res["total_score"], 77.5)
        self.assertEqual(score_res["status"], "Review")

    def test_candidate_ranking(self):
        """Test Python deterministic ranking by score."""
        candidates = [
            {"candidate_name": "Charlie", "total_score": 72.5, "experience": 3.0},
            {"candidate_name": "Alice", "total_score": 91.5, "experience": 5.0},
            {"candidate_name": "Bob", "total_score": 84.0, "experience": 4.0},
            {"candidate_name": "David", "total_score": 58.0, "experience": 1.0},
        ]
        ranked = rank_candidates(candidates, sort_by="score", ascending=False)
        self.assertEqual(ranked[0]["candidate_name"], "Alice")
        self.assertEqual(ranked[0]["rank"], 1)
        self.assertEqual(ranked[1]["candidate_name"], "Bob")
        self.assertEqual(ranked[2]["candidate_name"], "Charlie")
        self.assertEqual(ranked[3]["candidate_name"], "David")


if __name__ == "__main__":
    unittest.main()
