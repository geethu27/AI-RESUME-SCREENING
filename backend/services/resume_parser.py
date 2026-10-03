import os
import logging
from typing import Dict, Any, Tuple
import PyPDF2
import docx

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_path: str) -> Tuple[str, bool, str]:
    """
    Extracts raw text from a PDF file using PyPDF2.
    Rule-based extraction only. No OCR or AI models are used.

    Returns:
        (extracted_text, success, error_message)
    """
    if not os.path.exists(file_path):
        return "", False, f"File does not exist: {file_path}"

    text_pieces = []
    try:
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            if reader.is_encrypted:
                try:
                    # Attempt empty password decrypt
                    reader.decrypt("")
                except Exception:
                    return "", False, "PDF is password-protected and cannot be read."

            total_pages = len(reader.pages)
            if total_pages == 0:
                return "", False, "The PDF file contains no pages."

            for page_num in range(total_pages):
                page = reader.pages[page_num]
                page_text = page.extract_text() or ""
                if page_text.strip():
                    text_pieces.append(page_text)

        full_text = "\n".join(text_pieces).strip()

        # If full text is empty, it may be a scanned image-only PDF
        if not full_text:
            return (
                "",
                False,
                "Text could not be extracted from this PDF. It may contain scanned images without selectable text."
            )

        return full_text, True, ""

    except PyPDF2.errors.PdfReadError as e:
        logger.warning(f"PyPDF2 read error on {file_path}: {e}")
        return "", False, f"Invalid or corrupted PDF file: {str(e)}"
    except Exception as e:
        logger.error(f"Unexpected error extracting PDF {file_path}: {e}")
        return "", False, f"Error processing PDF file: {str(e)}"


def extract_text_from_docx(file_path: str) -> Tuple[str, bool, str]:
    """
    Extracts raw text from a DOCX file using python-docx.
    Rule-based extraction from paragraphs and tables. No AI/ML models are used.

    Returns:
        (extracted_text, success, error_message)
    """
    if not os.path.exists(file_path):
        return "", False, f"File does not exist: {file_path}"

    text_pieces = []
    try:
        doc = docx.Document(file_path)

        # Extract paragraphs
        for paragraph in doc.paragraphs:
            content = paragraph.text.strip()
            if content:
                text_pieces.append(content)

        # Extract text from tables if any
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_pieces.append(" | ".join(row_cells))

        full_text = "\n".join(text_pieces).strip()

        if not full_text:
            return "", False, "The DOCX file does not contain readable text content."

        return full_text, True, ""

    except Exception as e:
        logger.error(f"Error extracting DOCX {file_path}: {e}")
        return "", False, f"Invalid or corrupted DOCX file: {str(e)}"


def parse_resume(file_path: str) -> Dict[str, Any]:
    """
    Main resume parsing function.
    Detects file extension and routes to appropriate rule-based extractor.
    
    Returns:
        dict: {
            "text": str,
            "success": bool,
            "error": str or None,
            "filename": str,
            "file_type": str
        }
    """
    filename = os.path.basename(file_path)
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if extension == "pdf":
        text, success, error_msg = extract_text_from_pdf(file_path)
        return {
            "text": text,
            "success": success,
            "error": error_msg if not success else None,
            "filename": filename,
            "file_type": "pdf"
        }
    elif extension == "docx":
        text, success, error_msg = extract_text_from_docx(file_path)
        return {
            "text": text,
            "success": success,
            "error": error_msg if not success else None,
            "filename": filename,
            "file_type": "docx"
        }
    else:
        return {
            "text": "",
            "success": False,
            "error": f"Unsupported file format '.{extension}'. Only PDF and DOCX files are allowed.",
            "filename": filename,
            "file_type": extension
        }
